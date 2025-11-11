#!/usr/bin/env python3
"""
Detailed analysis of product section in CD.xlsx template
"""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def analyze_product_section():
    template_path = 'resources/sample/2/CD.xlsx'
    wb = load_workbook(template_path)
    ws = wb.active

    print("=" * 100)
    print("DETAILED ANALYSIS OF PRODUCT SECTION (Rows 160-220)")
    print("=" * 100)

    # Print all cells for rows 160-220
    for row_idx in range(160, min(221, ws.max_row + 1)):
        print(f"\n{'='*100}")
        print(f"ROW {row_idx}")
        print(f"{'='*100}")

        for col_idx in range(1, min(35, ws.max_column + 1)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                col_letter = get_column_letter(col_idx)
                cell_ref = f"{col_letter}{row_idx}"

                # Get formatting info
                number_format = cell.number_format if cell.number_format else 'General'

                value = cell.value
                # Handle long strings
                if isinstance(value, str) and len(value) > 80:
                    value_str = value[:80] + "..."
                else:
                    value_str = str(value)

                print(f"  {cell_ref:6} = {value_str:60} [format: {number_format}]")

if __name__ == "__main__":
    analyze_product_section()
