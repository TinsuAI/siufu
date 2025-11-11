#!/usr/bin/env python3
"""
Script to analyze CD.xlsx template structure for Story 2.5
This will help document cell mappings, formats, and structure
"""


from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def analyze_template():
    template_path = 'resources/sample/2/CD.xlsx'

    print(f"Loading template: {template_path}\n")
    wb = load_workbook(template_path)
    ws = wb.active

    print("=" * 80)
    print("TEMPLATE ANALYSIS: CD.xlsx")
    print("=" * 80)
    print(f"\nWorksheet name: {ws.title}")
    print(f"Max row: {ws.max_row}")
    print(f"Max column: {ws.max_column}")

    # Analyze merged cells
    print("\n" + "=" * 80)
    print("MERGED CELLS")
    print("=" * 80)
    for merged_range in ws.merged_cells.ranges:
        print(f"  {merged_range}")

    # Print all cell values and formatting for first 50 rows
    print("\n" + "=" * 80)
    print("CELL CONTENTS (First 50 rows)")
    print("=" * 80)

    for row_idx in range(1, min(51, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                col_letter = get_column_letter(col_idx)
                cell_ref = f"{col_letter}{row_idx}"

                # Get formatting info
                number_format = cell.number_format if cell.number_format else 'General'

                value_str = str(cell.value)[:50]  # Truncate long values
                row_data.append(f"{cell_ref}: '{value_str}' [fmt: {number_format}]")

        if row_data:
            print(f"\nRow {row_idx}:")
            for item in row_data:
                print(f"  {item}")

    # Column width analysis
    print("\n" + "=" * 80)
    print("COLUMN WIDTHS")
    print("=" * 80)
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        width = ws.column_dimensions[col_letter].width
        if width:
            print(f"  Column {col_letter}: {width}")

    # Try to identify key sections
    print("\n" + "=" * 80)
    print("IDENTIFYING KEY SECTIONS")
    print("=" * 80)

    # Look for Vietnamese keywords
    keywords = {
        'header': ['tên', 'địa chỉ', 'mã số thuế', 'ngày'],
        'product': ['hàng', 'mã hs', 'số lượng', 'đơn giá', 'thành tiền', 'xuất xứ'],
        'tax': ['thuế', 'vat', 'tổng']
    }

    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value and isinstance(cell.value, str):
                cell_value_lower = cell.value.lower()
                for section, keyword_list in keywords.items():
                    for keyword in keyword_list:
                        if keyword in cell_value_lower:
                            col_letter = get_column_letter(col_idx)
                            print(f"  {section.upper()}: Found '{keyword}' at {col_letter}{row_idx}: {cell.value[:50]}")

if __name__ == "__main__":
    analyze_template()
